# models
from .models import (
    Fine, Fee, FeeCategory, FeeCycle,
    FeeDiscount, ClassSectionFees, FeeConfiguration,
    StudentDiscount, FeeHeadMapping, Session, FeeStatus , StudentDiscountThrough
)
from classform.models import ClassRoomStudent, ClassRoom
from form.models import StudentInfo, ParentInfo, CurrentAddress

# forms
from .forms import (
    FeeCategoryForm, FeeCycleForm,
    FeeConfigurationForm, FeeDiscountForm, StudentDiscountForm,
    FeeHeadMasterForm
)

# 3rd party
from openpyxl import load_workbook
from tempfile import NamedTemporaryFile
from datetime import date
from io import BytesIO
from django.http import HttpResponse
from django.template.loader import get_template
from xhtml2pdf import pisa

# django
from django.core.files import File
from django.views import generic as generic_views
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.urls import reverse_lazy
from django.contrib.auth.mixins import LoginRequiredMixin
from django.forms import inlineformset_factory
from django.db.models import Q, Sum

# Create your views here.


@login_required
def add_fine(request):
    """
    Add Fine for a student
    Input: Admission Number of the student
    """
    if request.method == "GET":
        category = request.POST.get("category")
        fine = request.POST.get("fine")
        add_number = request.POST.get("add_number")
        description = request.POST.get("description")
        if add_number:
            student = ClassRoomStudent.objects.get(
                student__admissionNumber=add_number)
            Fine.objects.create(fine=fine, category=category,
                                description=description, student=student)
            redirect('addFine')
    return render(request, 'fees/addFine.html')


class FeeCategoryCreateView(generic_views.CreateView, LoginRequiredMixin):
    template_name = 'fees/addFeeCategory.html'
    form_class = FeeCategoryForm

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees category created"
        )
        return super().get_success_url()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['object_list'] = FeeCategory.objects.all()
        return context


class FeeCategoryUpdateView(generic_views.UpdateView, LoginRequiredMixin):
    template_name = 'fees/addFeeCategory.html'
    form_class = FeeCategoryForm
    queryset = FeeCategory.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees category sucessfully updated"
        )
        return super().get_success_url()


class FeeCategoryDeleteView(generic_views.DeleteView, LoginRequiredMixin):
    template_name = 'fees/deleteFeeCategory.html'
    queryset = FeeCategory.objects.all()
    success_url = reverse_lazy('feeCategories')

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees category deleted"
        )
        return super().get_success_url()


class FeeCycleCreateView(generic_views.CreateView, LoginRequiredMixin):
    template_name = 'fees/addFeeCycle.html'
    form_class = FeeCycleForm

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees cycle created"
        )
        return super().get_success_url()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['object_list'] = FeeCycle.objects.all()
        return context


class FeeCycleUpdateView(generic_views.UpdateView, LoginRequiredMixin):
    template_name = 'fees/addFeeCycle.html'
    form_class = FeeCycleForm
    queryset = FeeCycle.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees cycle sucessfully updated"
        )
        return super().get_success_url()


class FeeCycleDeleteView(generic_views.DeleteView, LoginRequiredMixin):
    template_name = 'fees/deleteFeeCycle.html'
    queryset = FeeCycle.objects.all()
    success_url = reverse_lazy('feeCycles')

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees cycle deleted"
        )
        return super().get_success_url()


class FeeConfigurationCreateView(generic_views.CreateView, LoginRequiredMixin):
    template_name = 'fees/addFeeConfiguration.html'
    form_class = FeeConfigurationForm

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees configuration created"
        )
        return super().get_success_url()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['object_list'] = FeeConfiguration.objects.all()
        return context


class FeeConfigurationDeleteView(generic_views.DeleteView, LoginRequiredMixin):
    template_name = 'fees/deleteFeeConfiguration.html'
    success_url = reverse_lazy('feeConfigurations')
    queryset = FeeConfiguration.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees Configuration deleted"
        )
        return super().get_success_url()


class FeeConfigurationUpdateView(generic_views.UpdateView, LoginRequiredMixin):
    template_name = 'fees/addFeeConfiguration.html'
    form_class = FeeConfigurationForm
    queryset = FeeConfiguration.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees configuration updated"
        )
        return super().get_success_url()


class FeeDiscountCreateView(generic_views.CreateView, LoginRequiredMixin):
    template_name = 'fees/addFeeDiscount.html'
    form_class = FeeDiscountForm

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees discount created"
        )
        return super().get_success_url()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['object_list'] = FeeDiscount.objects.all()
        return context


class FeeDiscountDeleteView(generic_views.DeleteView, LoginRequiredMixin):
    template_name = 'fees/deleteFeeDiscount.html'
    success_url = reverse_lazy('feeDiscounts')
    queryset = FeeDiscount.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees discount deleted"
        )
        return super().get_success_url()


class FeeDiscountUpdateView(generic_views.UpdateView, LoginRequiredMixin):
    template_name = 'fees/addFeeDiscount.html'
    form_class = FeeDiscountForm
    queryset = FeeDiscount.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Fees discount updated"
        )
        return super().get_success_url()


class StudentDiscountCreateView(generic_views.CreateView, LoginRequiredMixin):
    template_name = 'fees/addStudentDiscount.html'
    form_class = StudentDiscountForm

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Student discount created"
        )
        return super().get_success_url()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['object_list'] = StudentDiscount.objects.all()
        return context


class StudentDiscountDeleteView(generic_views.DeleteView, LoginRequiredMixin):
    template_name = 'fees/deleteStudentDiscount.html'
    success_url = reverse_lazy('studentDiscounts')
    queryset = StudentDiscount.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Student discount deleted"
        )
        return super().get_success_url()


class StudentDiscountUpdateView(generic_views.UpdateView, LoginRequiredMixin):
    template_name = 'fees/addStudentDiscount.html'
    form_class = StudentDiscountForm
    queryset = StudentDiscount.objects.all()

    def get_success_url(self):
        messages.add_message(
            self.request,
            messages.SUCCESS,
            message="Student discount updated"
        )
        return super().get_success_url()


class FeeHeadMappingFilter(generic_views.TemplateView, LoginRequiredMixin):
    template_name = 'fees/fee_head_mapping_filter.html'

    def post(self, request, *args, **kwargs):
        if request.POST.get('admission-no'):
            admission_no = int(request.POST.get('admission-no'))
        else:
            admission_no = int(request.POST.get('student', 0))

        return redirect(reverse_lazy('feeHeadMapping', kwargs={'admission_no': admission_no}))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["students"] = StudentInfo.objects.all()
        context["classrooms"] = ClassRoom.objects.all()
        return context


class FeeHeadMappingView(generic_views.FormView, LoginRequiredMixin):
    template_name = 'fees/fee_head_mapping.html'
    form_class = FeeHeadMasterForm

    def get_success_url(self):
        return self.request.path

    def get_object_list(self, student: StudentInfo):
        return FeeHeadMapping.objects.filter(student=student)

    def get_form_kwargs(self):
        addNo = self.kwargs.get('admission_no')
        kwargs = super().get_form_kwargs()
        kwargs['student_instance'] = StudentInfo.objects.get(
            admissionNumber=addNo)
        kwargs['month'] = int(self.request.GET.get('month', 0))
        return kwargs

    def form_valid(self, form):
        form.save()
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["students"] = StudentInfo.objects.all()
        context["classrooms"] = ClassRoom.objects.all()
        context['fee_categories'] = FeeCategory.objects.all()
        return context


@ login_required
def add_fee_class(request):
    """
    Set fee amount for a particular class
    Input: Fee amount, Class-Section and Category
    """
    if request.method == 'POST':
        fees = request.POST.get("fees")
        classSection = request.POST.get("class_room")
        category = request.POST.get("category")
        if category:
            fee_category = FeeCategory.objects.get(submission_type=category)
            classRoom = ClassRoom.objects.get(classSection=classSection)
            ClassSectionFees.objects.create(classSection=classRoom,
                                            feeCategory=fee_category, fees=fees)
            return redirect('addFeeClass')
    return render(request, 'fees/addClassFee.html', {"category": FeeCategory.objects.all(), "class_rooms": ClassRoom.objects.all()})


@ login_required
def fee_collection(request):
    """
    Get fee payment for a particular student and generate a fee slip.
    Input: Class, Amount, Payment Method, Admission Number, Months paid for and fee slip template
    Issue: Fee slip is generated but not exported to user in pdf.
    """
    if request.method == 'POST':
        classSection = request.POST.get("class_room")
        payment_method = request.POST.get("payment_method")
        amount = request.POST.get("amount")
        reg_no = request.POST.get("register_number")
        months = request.POST.get("months")
        add_number = request.POST.get("add_number")
        if "excel" in request.FILES:
            fee_slip_excell = request.FILES["excel"]
        if add_number and fee_slip_excell:
            classRoom = ClassRoom.objects.get(classSection=classSection)
            student = ClassRoomStudent.objects.get(
                student__admissionNumber=add_number)
            wb = load_workbook(fee_slip_excell)
            sheet = wb.get_sheet_by_name('Sheet2')
            # set slip no, name, class, date, months, class, amount

            # 1st copy
            sheet["B9"] = reg_no
            sheet["B10"] = add_number
            sheet["B12"] = student.student.admissionNumber
            sheet["G10"] = date.today().strftime("%B %d, %Y")
            sheet["G11"] = months
            sheet["E15"] = amount
            sheet["C20"] = amount

            # 2nd copy
            sheet["I9"] = reg_no
            sheet["I10"] = add_number
            sheet["I12"] = student.student.admissionNumber
            sheet["N10"] = date.today().strftime("%B %d, %Y")
            sheet["N11"] = months
            sheet["L15"] = amount
            sheet["J20"] = amount

            # wb.save('fee-slip.xlsx')
            with NamedTemporaryFile() as tmp:
                wb.save(tmp.name)
                tmp.seek(0)

                Fee.objects.create(regNo=reg_no, student=student, classSection=classRoom, submissionDate=date.today(
                ), monthsPaid=months, payment_method=payment_method, amount=amount)
                fee = Fee.objects.get(
                    student=student, regNo=reg_no, submissionDate=date.today())
                fee.feeSlip.save("fee-slip", File(tmp), True)
                fee.save()

                try:
                    from win32com import client
                    import win32api

                    excel = client.DispatchEx("Excel.Application")
                    excel.Visible = 0

                    wb = excel.Workbooks.Open(tmp)
                    ws = wb.Worksheets[1]
                    try:
                        wb.SaveAs('c:\\targetfolder\\feeslip.pdf',
                                  FileFormat=57)
                    except:
                        print("Failed to convert")
                    finally:
                        wb.Close()
                        excel.Quit()
                except:
                    pass

    return render(request, 'fees/collectFee.html', {"class_rooms": ClassRoom.objects.all()})


def render_to_pdf(template_src, context_dict={}):
    template = get_template(template_src)
    html = template.render(context_dict)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("ISO-8859-1")), result)
    if not pdf.err:
        return HttpResponse(result.getvalue(), content_type='application/pdf')
    return None


# def fee_slip(request):
#     stu = StudentInfo.objects.get(admissionNumber=1)
#     # print(stu.admissionNumber)
#     par = ParentInfo.objects.get(student=stu.admissionNumber)
#     data = {
#         'stu': stu,
#         'par': par,
#         'pagesize': 'A4',
#     }
#     # pdf = render_to_pdf('fees/fee_slip.html', data)
#     # return HttpResponse(pdf, content_type='application/pdf')
#     return render(request,'fees/fee_slip.html',data)


def fee_slip(request,admission_no,session_id):
    stu = StudentInfo.objects.get(admissionNumber=admission_no)
    par = ParentInfo.objects.get(student=stu.admissionNumber)
    session = Session.objects.get(pk=session_id)
    data = {
        'stu': stu,
        'par': par,
        'session':session,
    }
    template = get_template("fees/fee_slip.html")
    data_p = template.render(data)
    response = BytesIO()

    pdfPage = pisa.pisaDocument(BytesIO(data_p.encode("UTF-8")), response)
    if not pdfPage.err:
        return HttpResponse(response.getvalue(), content_type="application/pdf")
    else:
        return HttpResponse("Error Generating PDF")

# @login_required
# def fee_collect(request):
#     context = {'Session':Session.objects.all(), 'class_rooms':ClassRoom.objects.all(), 'student':ClassRoomStudent.objects.all()}
#     return render(request,'fees/fee_collection.html',context)


class FeeCollectionFilter(generic_views.TemplateView, LoginRequiredMixin):
    template_name = 'fees/fee_collection_filter.html'

    def post(self, request, *args, **kwargs):
        if request.POST.get('admission-no'):
            admission_no = int(request.POST.get('admission-no'))
        else:
            admission_no = int(request.POST.get('student', 0))

        if request.POST.get('session'):
            session_id = int(request.POST.get('session'))
        else:
            session_id = int(request.POST.get('session', 0))

        return redirect(reverse_lazy('feeCollection', kwargs={'admission_no': admission_no, 'session_id': session_id}))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["students"] = StudentInfo.objects.all()
        context["classrooms"] = ClassRoom.objects.all()
        context["Session"] = Session.objects.all()
        return context

fee_dict={'amount':0 , 'remarks':'Paid','final_amt':0}
class FeeCollection(generic_views.TemplateView, LoginRequiredMixin):
    template_name = 'fees/fee_collection.html'

    def get_success_url(self):
        return self.request.path

    def get_fee_cycles(self):
        fee_cycles = []
        admission_no = self.kwargs.get('admission_no')
        student = StudentInfo.objects.get(admissionNumber=admission_no)
        for month_range in self.request.GET.getlist('month'):
            month_range = list(map(int, month_range.split('-')))
            if len(month_range) == 1:
                month_range.append(month_range[0])
            month_range = range(month_range[0], month_range[1]+1)
            fee_head_mapping = student.fee_head_mappings.all().filter(month__in=month_range)
            total_amount = fee_head_mapping.aggregate(Sum('amount'))[
                'amount__sum']
            fee_cycles.append((fee_head_mapping, total_amount))
        return fee_cycles

    def get_student_discount(self):
        admission_no = self.kwargs.get('admission_no')
        student = StudentInfo.objects.get(admissionNumber=admission_no)
        discount = StudentDiscountThrough.objects.filter(student=admission_no)
        return discount

    def get_final_amt(self):
        discount = self.get_student_discount()
        total = self.get_fee_cycles()
        bal = self.get_balance()
        for i in discount:
            if not total:
                fee_dict['final_amt'] = i.amount
            else:
                fee_dict['final_amt'] = total[0][1] - i.amount
        # print(discount.amount)
        if not bal:
            grand_total = fee_dict.get('final_amt')
        else:
            grand_total = fee_dict.get('final_amt') + bal[0]
        return grand_total


    def get_balance(self):
        balances = []
        admission_no = self.kwargs.get('admission_no')
        student = StudentInfo.objects.get(admissionNumber=admission_no)
        session_id = self.kwargs.get('session_id')
        for month in self.request.GET.getlist('month'):
            balance = 0
            try:
                status = FeeStatus.objects.filter(
                    session__pk=session_id,
                    student=student,
                    month__lte=month
                )
                balance += status.aggregate(Sum('balance'))['balance__sum']
            except:
                balance += 0
            balances.append(balance)

        return balances

    def post(self, request, *args, **kwargs):
        fee_cycles = self.get_fee_cycles()
        session = Session.objects.get(pk=self.kwargs.get('session_id'))
        addno = self.kwargs.get('admission_no')
        session_id = self.kwargs.get('session_id')
        sub_date = self.request.POST.get('sub_date','')
        remarks = self.request.POST.get('remarks','')
        pay_method = self.request.POST.get('session','')
        # classSection = ClassRoom.objects.filter(classSection=StudentInfo.objects.filter(admissionNumber=addno).values_list('classSection')).values_list('classSection')
        id_data = ClassRoomStudent.objects.only('classRoom').get(student=addno).id
        classSection= ClassRoom.objects.only('id').get(id=id_data)
        student = ClassRoomStudent.objects.get(student=addno)
        amount = self.request.POST.get('deposit','')
        fee_dict['amount'] = amount
        fee_dict['remarks'] = remarks
        months_paid = self.request.GET.getlist('month','')
        months_paid = list(map(int,months_paid))
        reg_no = addno
        fee = Fee(regNo=reg_no,payment_method=pay_method,submissionDate=sub_date,amount=amount,monthsPaid=months_paid[0],classSection=classSection,student=student)
        fee.save()
        m = self.auto_fee_slip()
        for index, (fees, total_amount) in enumerate(fee_cycles):
            pass
            amount_paid = request.POST.getlist(f'deposit')[index]
            amount_paid = 0 if amount_paid == '' else int(amount_paid)
            for fee in fees:
                try:
                    fee_status = FeeStatus.objects.get(
                        student=fee.student, session=session, month=fee.month)
                except:
                    fee_status = FeeStatus.objects.create(
                        student=fee.student, session=session,
                        month=fee.month, total_amount=0,
                        received=0, balance=0
                    )
                fee_status.total_amount = total_amount
                fee_status.received = amount_paid
                fee_status.balance = total_amount-amount_paid
                fee_status.save()
        return m

    def auto_fee_slip(self):
        addno = self.kwargs.get('admission_no')
        session_id = self.kwargs.get('session_id')
        stu = StudentInfo.objects.get(admissionNumber=addno)
        par = ParentInfo.objects.get(student=stu.admissionNumber)
        session = Session.objects.get(pk=session_id)
        fee_cycles = self.get_fee_cycles()
        balances = self.get_balance()
        disc = self.get_student_discount()
        amt = self.get_final_amt()
        fee_status = FeeStatus.objects.get(student=addno)
        data = {
            'stu': stu,
            'par': par,
            'session':session,
            'fee_cycle':fee_cycles,
            'balance':balances,
            'discount':disc,
            'final_amt':amt,
            'fee_status':fee_status,
            'amount':fee_dict.get('amount'),
            'remarks':fee_dict.get('remarks')
            }
        template = get_template("fees/fee_slip.html")
        data_p = template.render(data)
        response = BytesIO()

        pdfPage = pisa.pisaDocument(BytesIO(data_p.encode("UTF-8")), response)
        if not pdfPage.err:
            return HttpResponse(response.getvalue(), content_type="application/pdf")
        else:
            return HttpResponse("Error Generating PDF")


    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        addno = self.kwargs.get('admission_no')
        session_id = self.kwargs.get('session_id')
        # self.get_success_url()
        context["student"] = StudentInfo.objects.filter(admissionNumber=addno)
        context["parent_info"] = ParentInfo.objects.filter(student=addno)
        context["address"] = CurrentAddress.objects.filter(student=addno)
        context["students"] = StudentInfo.objects.all()
        context["classrooms"] = ClassRoom.objects.all()
        context["Session"] = Session.objects.all()
        context['cycle'] = FeeCycle.objects.filter(session=session_id)
        fee_cycles = self.get_fee_cycles()
        balances = self.get_balance()
        disc = self.get_student_discount()
        context['grand_total'] = self.get_final_amt()
        context['fee_details'] = zip(fee_cycles, balances,disc)
        return context
